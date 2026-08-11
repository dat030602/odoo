# -*- coding: utf-8 -*-
"""
test_template_engine.py
=======================

Unit tests for the TemplateEngine class.

These tests run WITHOUT an Odoo database — they use plain openpyxl
workbooks and dict contexts, making them fast and isolated.

Run with:
    pytest tests/test_template_engine.py -v --import-mode=importlib --rootdir=tests --confcutdir=tests
"""

import pytest
import openpyxl
import importlib.util
import os

# Import directly from the module file to avoid triggering the package
# __init__.py chain (which imports Odoo-dependent modules).
_spec = importlib.util.spec_from_file_location(
    "template_engine",
    os.path.join(os.path.dirname(__file__), "..", "models", "template_engine.py"),
)
_template_engine_module = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_template_engine_module)

TemplateEngine = _template_engine_module.TemplateEngine
TemplateError = _template_engine_module.TemplateError


# ─────────────────────────────────────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def make_sheet(rows):
    """Create a worksheet from a list of row lists.

    Args:
        rows: List of lists, where each inner list represents a row.
              Each element is a cell value (str, int, float, etc.).

    Returns:
        openpyxl Worksheet with the given data.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, row_values in enumerate(rows, start=1):
        for c, value in enumerate(row_values, start=1):
            ws.cell(row=r, column=c, value=value)
    return ws


# ─────────────────────────────────────────────────────────────────────────────
# TESTS: Placeholder resolution
# ─────────────────────────────────────────────────────────────────────────────

class TestPlaceholderResolution:
    """Tests for {{ }} placeholder substitution."""

    def test_simple_placeholder(self):
        """A single {{ key }} is replaced with the context value."""
        ws = make_sheet([['Hello {{ name }}']])
        engine = TemplateEngine(ws)
        engine.render({'name': 'World'})
        assert ws.cell(row=1, column=1).value == 'Hello World'

    def test_dot_path_resolution(self):
        """Dot-paths resolve through nested dicts."""
        ws = make_sheet([['{{ company.name }}']])
        engine = TemplateEngine(ws)
        engine.render({'company': {'name': 'Acme Corp'}})
        assert ws.cell(row=1, column=1).value == 'Acme Corp'

    def test_multiple_placeholders_in_one_cell(self):
        """Multiple placeholders in a single cell are all replaced."""
        ws = make_sheet([['{{ greeting }}, {{ name }}!']])
        engine = TemplateEngine(ws)
        engine.render({'greeting': 'Hello', 'name': 'Alice'})
        assert ws.cell(row=1, column=1).value == 'Hello, Alice!'

    def test_missing_key_resolves_to_empty(self):
        """A missing key resolves to empty string (not an error)."""
        ws = make_sheet([['{{ missing }}']])
        engine = TemplateEngine(ws)
        engine.render({})
        assert ws.cell(row=1, column=1).value == ''

    def test_none_value_resolves_to_empty(self):
        """A None value resolves to empty string."""
        ws = make_sheet([['{{ value }}']])
        engine = TemplateEngine(ws)
        engine.render({'value': None})
        assert ws.cell(row=1, column=1).value == ''

    def test_fmt_filter_sets_number_format(self):
        """The | fmt filter sets cell.number_format and preserves value type."""
        ws = make_sheet([['{{ amount | fmt:#,##0.00 }}']])
        engine = TemplateEngine(ws)
        engine.render({'amount': 1234.5})
        cell = ws.cell(row=1, column=1)
        assert cell.value == 1234.5
        assert cell.number_format == '#,##0.00'

    def test_formula_value_preserved(self):
        """A value starting with = is written as a formula."""
        ws = make_sheet([['{{ formula }}']])
        engine = TemplateEngine(ws)
        engine.render({'formula': '=SUM(A1:A10)'})
        assert ws.cell(row=1, column=1).value == '=SUM(A1:A10)'


# ─────────────────────────────────────────────────────────────────────────────
# TESTS: For loop blocks
# ─────────────────────────────────────────────────────────────────────────────

class TestForBlocks:
    """Tests for {% for %} ... {% endfor %} loop expansion."""

    def test_flat_for_expands_correctly(self):
        """A flat for loop expands one row per item."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{{ line.name }}', '{{ line.amount }}'],
            ['{% endfor %}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({
            'lines': [
                {'name': 'A', 'amount': 10},
                {'name': 'B', 'amount': 20},
            ],
        })
        assert ws.cell(row=1, column=1).value == 'A'
        assert ws.cell(row=1, column=2).value == 10
        assert ws.cell(row=2, column=1).value == 'B'
        assert ws.cell(row=2, column=2).value == 20

    def test_for_block_marker_rows_deleted(self):
        """After expansion, the {% for %} and {% endfor %} marker rows are removed."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{{ line.name }}'],
            ['{% endfor %}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({'lines': [{'name': 'A'}, {'name': 'B'}]})
        # Only 2 rows should remain (one per item)
        assert ws.max_row == 2
        assert ws.cell(row=1, column=1).value == 'A'
        assert ws.cell(row=2, column=1).value == 'B'

    def test_empty_list_deletes_block(self):
        """An empty list deletes the entire block (markers + body)."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{{ line.name }}'],
            ['{% endfor %}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({'lines': []})
        # All 3 rows should be deleted
        assert ws.cell(row=1, column=1).value is None

    def test_single_item_for_loop(self):
        """A single-item list expands to one row."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{{ line.name }}'],
            ['{% endfor %}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({'lines': [{'name': 'Only'}]})
        assert ws.cell(row=1, column=1).value == 'Only'
        assert ws.cell(row=2, column=1).value is None

    def test_multi_row_body_per_item(self):
        """A loop body spanning multiple rows is duplicated per item."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{{ line.name }}'],
            ['Detail: {{ line.detail }}'],
            ['{% endfor %}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({
            'lines': [
                {'name': 'A', 'detail': 'A-detail'},
                {'name': 'B', 'detail': 'B-detail'},
            ],
        })
        assert ws.cell(row=1, column=1).value == 'A'
        assert ws.cell(row=2, column=1).value == 'Detail: A-detail'
        assert ws.cell(row=3, column=1).value == 'B'
        assert ws.cell(row=4, column=1).value == 'Detail: B-detail'

    def test_for_block_with_surrounding_rows(self):
        """Rows above and below the for block are preserved."""
        ws = make_sheet([
            ['Header'],
            ['{% for line in lines %}'],
            ['{{ line.name }}'],
            ['{% endfor %}'],
            ['Footer'],
        ])
        engine = TemplateEngine(ws)
        engine.render({'lines': [{'name': 'A'}, {'name': 'B'}]})
        assert ws.cell(row=1, column=1).value == 'Header'
        assert ws.cell(row=2, column=1).value == 'A'
        assert ws.cell(row=3, column=1).value == 'B'
        assert ws.cell(row=4, column=1).value == 'Footer'


# ─────────────────────────────────────────────────────────────────────────────
# TESTS: If blocks
# ─────────────────────────────────────────────────────────────────────────────

class TestIfBlocks:
    """Tests for {% if %} ... {% endif %} conditional blocks."""

    def test_if_true_keeps_body(self):
        """When the condition is True, the body is kept and markers removed."""
        ws = make_sheet([
            ['{% if show %}'],
            ['Visible content'],
            ['{% endif %}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({'show': True})
        assert ws.cell(row=1, column=1).value == 'Visible content'
        assert ws.cell(row=2, column=1).value is None

    def test_if_false_removes_body(self):
        """When the condition is False, the entire block is removed."""
        ws = make_sheet([
            ['{% if show %}'],
            ['Visible content'],
            ['{% endif %}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({'show': False})
        assert ws.cell(row=1, column=1).value is None

    def test_if_with_comparison(self):
        """Conditions with comparison operators work correctly."""
        ws = make_sheet([
            ['{% if amount > 100 %}'],
            ['Big'],
            ['{% endif %}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({'amount': 200})
        assert ws.cell(row=1, column=1).value == 'Big'

    def test_if_with_comparison_false(self):
        """Conditions with comparison operators return False correctly."""
        ws = make_sheet([
            ['{% if amount > 100 %}'],
            ['Big'],
            ['{% endif %}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({'amount': 50})
        assert ws.cell(row=1, column=1).value is None


# ─────────────────────────────────────────────────────────────────────────────
# TESTS: Aggregate filters (sum, count, avg, max, min)
# ─────────────────────────────────────────────────────────────────────────────

class TestAggregateFilters:
    """Tests for {{ list | sum:'field' }} and related filters."""

    def test_sum_filter_produces_subtotal_formula(self):
        """The sum filter generates a SUBTOTAL(9, ...) formula."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{{ line.amount }}'],
            ['{% endfor %}'],
            ['{{ lines | sum:\'amount\' }}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({
            'lines': [
                {'amount': 10},
                {'amount': 20},
                {'amount': 30},
            ],
        })
        # 3 items → rows 1-3, sum formula on row 4
        assert ws.cell(row=4, column=1).value == '=SUBTOTAL(9,A1:A3)'

    def test_count_filter_produces_subtotal_formula(self):
        """The count filter generates a SUBTOTAL(2, ...) formula."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{{ line.amount }}'],
            ['{% endfor %}'],
            ['{{ lines | count:\'amount\' }}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({
            'lines': [
                {'amount': 10},
                {'amount': 20},
                {'amount': 30},
            ],
        })
        assert ws.cell(row=4, column=1).value == '=SUBTOTAL(2,A1:A3)'

    def test_avg_filter_produces_subtotal_formula(self):
        """The avg filter generates a SUBTOTAL(1, ...) formula."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{{ line.amount }}'],
            ['{% endfor %}'],
            ['{{ lines | avg:\'amount\' }}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({
            'lines': [
                {'amount': 10},
                {'amount': 20},
                {'amount': 30},
            ],
        })
        assert ws.cell(row=4, column=1).value == '=SUBTOTAL(1,A1:A3)'

    def test_max_filter_produces_subtotal_formula(self):
        """The max filter generates a SUBTOTAL(4, ...) formula."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{{ line.amount }}'],
            ['{% endfor %}'],
            ['{{ lines | max:\'amount\' }}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({
            'lines': [
                {'amount': 10},
                {'amount': 20},
                {'amount': 30},
            ],
        })
        assert ws.cell(row=4, column=1).value == '=SUBTOTAL(4,A1:A3)'

    def test_min_filter_produces_subtotal_formula(self):
        """The min filter generates a SUBTOTAL(5, ...) formula."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{{ line.amount }}'],
            ['{% endfor %}'],
            ['{{ lines | min:\'amount\' }}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({
            'lines': [
                {'amount': 10},
                {'amount': 20},
                {'amount': 30},
            ],
        })
        assert ws.cell(row=4, column=1).value == '=SUBTOTAL(5,A1:A3)'

    def test_empty_list_sum_is_zero(self):
        """When the loop renders zero rows, the sum formula is 0."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{{ line.amount }}'],
            ['{% endfor %}'],
            ['{{ lines | sum:\'amount\' }}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({'lines': []})
        assert ws.cell(row=1, column=1).value == 0

    def test_aggregate_filter_unknown_var_raises(self):
        """An aggregate filter referencing an unknown loop variable raises TemplateError."""
        ws = make_sheet([
            ['{{ unknown | sum:\'amount\' }}'],
        ])
        engine = TemplateEngine(ws)
        with pytest.raises(TemplateError, match="not a known expanded loop variable"):
            engine.render({'lines': []})


# ─────────────────────────────────────────────────────────────────────────────
# TESTS: Error handling
# ─────────────────────────────────────────────────────────────────────────────

class TestErrorHandling:
    """Tests for template syntax error detection."""

    def test_unclosed_for_block_raises(self):
        """An unclosed {% for %} raises TemplateError."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{{ line.name }}'],
        ])
        engine = TemplateEngine(ws)
        with pytest.raises(TemplateError, match="Unclosed"):
            engine.render({'lines': []})

    def test_unclosed_if_block_raises(self):
        """An unclosed {% if %} raises TemplateError."""
        ws = make_sheet([
            ['{% if show %}'],
            ['Content'],
        ])
        engine = TemplateEngine(ws)
        with pytest.raises(TemplateError, match="Unclosed"):
            engine.render({'show': True})

    def test_unmatched_endfor_raises(self):
        """An orphaned {% endfor %} raises TemplateError."""
        ws = make_sheet([
            ['Content'],
            ['{% endfor %}'],
        ])
        engine = TemplateEngine(ws)
        with pytest.raises(TemplateError, match="Unmatched"):
            engine.render({})

    def test_mismatched_block_raises(self):
        """A {% endif %} closing a {% for %} raises TemplateError."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{{ line.name }}'],
            ['{% endif %}'],
        ])
        engine = TemplateEngine(ws)
        with pytest.raises(TemplateError, match="Mismatched"):
            engine.render({'lines': []})

    def test_empty_loop_body_raises(self):
        """A for block with no body rows raises TemplateError."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{% endfor %}'],
        ])
        engine = TemplateEngine(ws)
        with pytest.raises(TemplateError, match="Empty loop body"):
            engine.render({'lines': [{'name': 'A'}]})

    def test_nested_loop_raises(self):
        """Nested for loops raise TemplateError in Giai đoạn 1."""
        ws = make_sheet([
            ['{% for g in groups %}'],
            ['{% for l in g.lines %}'],
            ['{{ l.name }}'],
            ['{% endfor %}'],
            ['{% endfor %}'],
        ])
        engine = TemplateEngine(ws)
        with pytest.raises(TemplateError, match="Nested for loops"):
            engine.render({'groups': []})


# ─────────────────────────────────────────────────────────────────────────────
# TESTS: Style preservation
# ─────────────────────────────────────────────────────────────────────────────

class TestStylePreservation:
    """Tests that row styles are copied during loop expansion."""

    def test_styles_copied_to_expanded_rows(self):
        """Styles from the template body row are copied to expanded rows."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = '{% for line in lines %}'
        ws['A2'] = '{{ line.name }}'
        ws['A3'] = '{% endfor %}'

        # Apply a style to the template body row
        from openpyxl.styles import Font
        ws['A2'].font = Font(bold=True, color='FF0000')

        engine = TemplateEngine(ws)
        engine.render({'lines': [{'name': 'A'}, {'name': 'B'}]})

        # Both expanded rows should have the bold font
        assert ws.cell(row=1, column=1).font.bold is True
        assert ws.cell(row=2, column=1).font.bold is True

    def test_row_height_preserved(self):
        """Row height from the template body is copied to expanded rows."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = '{% for line in lines %}'
        ws['A2'] = '{{ line.name }}'
        ws['A3'] = '{% endfor %}'
        ws.row_dimensions[2].height = 30.0

        engine = TemplateEngine(ws)
        engine.render({'lines': [{'name': 'A'}, {'name': 'B'}]})

        # After expansion and marker deletion, the data rows should have
        # the height from the template body row.
        # Note: openpyxl's delete_rows doesn't shift row_dimensions,
        # so we check the dimensions that were set during copy.
        assert ws.row_dimensions[2].height == 30.0
        assert ws.row_dimensions[3].height == 30.0


# ─────────────────────────────────────────────────────────────────────────────
# TESTS: Integration scenarios
# ─────────────────────────────────────────────────────────────────────────────

class TestIntegrationScenarios:
    """End-to-end template rendering scenarios."""

    def test_full_report_with_header_and_total(self):
        """A complete report with header, loop, and grand total."""
        ws = make_sheet([
            ['{{ company.name }}'],
            ['Report printed on {{ print_date }}'],
            ['Name', 'Amount'],
            ['{% for line in lines %}'],
            ['{{ line.name }}', '{{ line.amount | fmt:\'#,##0.00\' }}'],
            ['{% endfor %}'],
            ['Grand Total', '{{ lines | sum:\'amount\' }}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({
            'company': {'name': 'Acme Corp'},
            'print_date': '01/08/2026',
            'lines': [
                {'name': 'Order A', 'amount': 100.50},
                {'name': 'Order B', 'amount': 200.75},
                {'name': 'Order C', 'amount': 300.25},
            ],
        })

        # Header rows preserved
        assert ws.cell(row=1, column=1).value == 'Acme Corp'
        assert ws.cell(row=2, column=1).value == 'Report printed on 01/08/2026'
        assert ws.cell(row=3, column=1).value == 'Name'
        assert ws.cell(row=3, column=2).value == 'Amount'

        # Loop rows expanded (3 items)
        assert ws.cell(row=4, column=1).value == 'Order A'
        assert ws.cell(row=4, column=2).value == 100.50
        assert ws.cell(row=4, column=2).number_format == '#,##0.00'
        assert ws.cell(row=5, column=1).value == 'Order B'
        assert ws.cell(row=5, column=2).value == 200.75
        assert ws.cell(row=6, column=1).value == 'Order C'
        assert ws.cell(row=6, column=2).value == 300.25

        # Grand total formula
        assert ws.cell(row=7, column=1).value == 'Grand Total'
        assert ws.cell(row=7, column=2).value == '=SUBTOTAL(9,B4:B6)'

    def test_if_inside_for_loop(self):
        """An if block inside a for loop works correctly."""
        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{% if line.amount > 100 %}'],
            ['{{ line.name }} (VIP)'],
            ['{% endif %}'],
            ['{% endfor %}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({
            'lines': [
                {'name': 'A', 'amount': 50},
                {'name': 'B', 'amount': 200},
                {'name': 'C', 'amount': 300},
            ],
        })

        # Item A: amount=50, condition false → no row
        # Item B: amount=200, condition true → VIP row
        # Item C: amount=300, condition true → VIP row
        assert ws.cell(row=1, column=1).value == 'B (VIP)'
        assert ws.cell(row=2, column=1).value == 'C (VIP)'
        # No more data rows (openpyxl max_row may be stale after delete_rows)
        assert ws.cell(row=3, column=1).value is None


# ─────────────────────────────────────────────────────────────────────────────
# TESTS: Image filter
# ─────────────────────────────────────────────────────────────────────────────

class TestImageFilter:
    """Tests for {{ value | image }} filter."""

    def test_image_filter_with_string_value(self):
        """The image filter correctly handles string base64 values."""
        import base64 as b64
        # 1x1 transparent PNG
        png_b64 = 'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=='

        ws = make_sheet([['{{ photo | image }}']])
        engine = TemplateEngine(ws)
        engine.render({'photo': png_b64})

        cell = ws.cell(row=1, column=1)
        assert cell.value is None
        assert len(ws._images) == 1
        assert ws._images[0].anchor == 'A1'

    def test_image_filter_with_bytes_value(self):
        """The image filter correctly handles bytes values (e.g. Odoo binary fields).

        Previously, bytes values were converted to their repr (b'...') which
        broke base64 decoding. The fix decodes bytes to str before embedding.
        """
        import base64 as b64
        # 1x1 transparent PNG
        png_b64 = 'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=='

        ws = make_sheet([['{{ photo | image }}']])
        engine = TemplateEngine(ws)
        engine.render({'photo': png_b64.encode('utf-8')})

        cell = ws.cell(row=1, column=1)
        assert cell.value is None
        assert len(ws._images) == 1
        assert ws._images[0].anchor == 'A1'

    def test_image_filter_in_for_loop(self):
        """The image filter works correctly inside for loops."""
        import base64 as b64
        # 1x1 transparent PNG
        png_b64 = 'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=='

        ws = make_sheet([
            ['{% for line in lines %}'],
            ['{{ line.photo | image }}'],
            ['{% endfor %}'],
        ])
        engine = TemplateEngine(ws)
        engine.render({
            'lines': [
                {'photo': png_b64},
                {'photo': png_b64},
            ],
        })

        assert ws.cell(row=1, column=1).value is None
        assert ws.cell(row=2, column=1).value is None
        assert len(ws._images) == 2
        assert ws._images[0].anchor == 'A1'
        assert ws._images[1].anchor == 'A2'

    def test_image_filter_with_url(self):
        """The image filter handles URL values (marker remains if fetch fails)."""
        ws = make_sheet([['{{ photo | image }}']])
        engine = TemplateEngine(ws)
        engine.render({'photo': 'https://example.com/image.png'})

        cell = ws.cell(row=1, column=1)
        # Since the URL won't resolve, the image won't be embedded,
        # and the marker remains in the cell value
        assert cell.value == '__IMAGE__:https://example.com/image.png'
